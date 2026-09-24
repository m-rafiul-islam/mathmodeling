from pathlib import Path

from pptx import Presentation
from pptx.util import Pt


OUTPUT_FILE = Path(__file__).resolve().parent / "Mathematical_Modeling_Differential_Equations.pptx"

TITLE = "Mathematical Modeling with Differential Equations"
SUBTITLE = "Math 4370.01 – Math Modeling of Natural Processes"

SLIDES = [
    (
        "Learning Objectives",
        [
            "Understand what a mathematical model is",
            "Translate real-world processes into differential equations",
            "Classify ODEs (order, linearity, autonomous vs. non-autonomous)",
            "Interpret solutions qualitatively (equilibria, stability)",
            "Apply models to natural/physical processes",
        ],
    ),
    (
        "What Is a Mathematical Model?",
        [
            "A simplified mathematical description of a real system",
            "Steps: Observe -> Assume -> Formulate -> Solve -> Validate",
            "Differential equations model rates of change (population growth, cooling, motion, chemical reactions)",
        ],
    ),
    (
        "Why Differential Equations?",
        [
            "Natural processes change continuously over time/space",
            "DEs relate a quantity to its rate of change: dy/dt = f(t, y)",
            "Powerful tool for prediction and understanding mechanisms",
        ],
    ),
    (
        "Classifying Differential Equations",
        [
            "Order: highest derivative present (1st order, 2nd order...)",
            "Linear vs. Nonlinear",
            "Autonomous (dy/dt = f(y)) vs. Non-autonomous (dy/dt = f(t, y))",
            "ODE vs. PDE",
        ],
    ),
    (
        "Example 1: Exponential Growth/Decay",
        [
            "Model: dP/dt = kP",
            "Solution: P(t) = P0 e^(kt)",
            "Applications: population growth, radioactive decay, compound interest",
        ],
    ),
    (
        "Example 2: Logistic Growth",
        [
            "Model: dP/dt = rP(1 - P/K)",
            "Accounts for carrying capacity K",
            "Equilibria: P = 0 (unstable), P = K (stable)",
        ],
    ),
    (
        "Example 3: Newton's Law of Cooling",
        [
            "Model: dT/dt = -k(T - T_env)",
            "Solution: T(t) = T_env + (T0 - T_env) e^(-kt)",
            "Application: cooling coffee, forensic time-of-death estimation",
        ],
    ),
    (
        "Qualitative Analysis: Equilibria & Stability",
        [
            "Equilibrium solutions: where dy/dt = 0",
            "Stability via phase line analysis",
            "Sign of derivative tells direction of flow",
        ],
    ),
    (
        "Systems of Differential Equations",
        [
            "Predator-Prey (Lotka-Volterra) model: dx/dt = ax - bxy, dy/dt = -cy + dxy",
            "Introduces interacting quantities and equilibrium points",
        ],
    ),
    (
        "Modeling Process Recap",
        [
            "Define variables & parameters",
            "State assumptions",
            "Write governing equation(s)",
            "Solve analytically/numerically",
            "Interpret & validate against data",
        ],
    ),
    (
        "In-Class Activity / Discussion",
        [
            "Given a real process (e.g., spread of a rumor, drug concentration in blood), have students set up the DE",
            "Discuss assumptions and limitations",
        ],
    ),
    (
        "Summary",
        [
            "DEs are essential tools for modeling natural processes",
            "Key models: exponential, logistic, cooling, predator-prey",
            "Next class: solving techniques (separation of variables, numerical methods)",
        ],
    ),
    ("Questions?", []),
]


def _set_run_font(paragraph, size):
    for run in paragraph.runs:
        run.font.size = Pt(size)


def add_title_slide(presentation: Presentation) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = TITLE
    slide.placeholders[1].text = SUBTITLE

    _set_run_font(slide.shapes.title.text_frame.paragraphs[0], 28)
    _set_run_font(slide.placeholders[1].text_frame.paragraphs[0], 18)


def add_bullet_slide(presentation: Presentation, title: str, bullets: list[str]) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = title
    _set_run_font(slide.shapes.title.text_frame.paragraphs[0], 24)

    body = slide.placeholders[1].text_frame
    body.clear()

    if not bullets:
        paragraph = body.paragraphs[0]
        paragraph.text = ""
        return

    for index, bullet in enumerate(bullets):
        paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        _set_run_font(paragraph, 20)


def build_presentation() -> Presentation:
    presentation = Presentation()
    add_title_slide(presentation)
    for title, bullets in SLIDES:
        add_bullet_slide(presentation, title, bullets)
    return presentation


def main() -> None:
    presentation = build_presentation()
    presentation.save(OUTPUT_FILE)
    print(f"Saved presentation to {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
